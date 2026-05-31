"""
DOCX Resume Parser
Extracts text content from DOCX files
"""

import docx
from typing import Optional
import logging

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


def extract_docx(file_path: str) -> Optional[str]:
    """
    Extract text from DOCX resume
    
    Args:
        file_path (str): Path to DOCX file
        
    Returns:
        Optional[str]: Extracted text or None if error occurs
    """
    try:
        doc = docx.Document(file_path)
        text = []
        
        # Extract text from paragraphs
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text.append(paragraph.text)
        
        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text.append(cell.text)
        
        full_text = '\n'.join(text)
        
        if not full_text.strip():
            logger.warning(f"No text extracted from {file_path}")
            return None
            
        logger.info(f"Successfully extracted {len(full_text)} characters from DOCX")
        return full_text.strip()
        
    except FileNotFoundError:
        logger.error(f"File not found: {file_path}")
        return None
    except Exception as e:
        logger.error(f"Unexpected error extracting DOCX: {str(e)}")
        return None

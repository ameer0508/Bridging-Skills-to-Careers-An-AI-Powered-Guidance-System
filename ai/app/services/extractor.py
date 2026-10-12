import os
try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    import docx
except ImportError:
    docx = None


def extract_text_from_file(file_path: str) -> str:
    """
    Extracts raw text from a PDF, DOCX, or TXT file.
    Raises ValueError when the document cannot be read; it never returns placeholder text.
    """
    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.pdf':
        return extract_text_from_pdf(file_path)
    if ext == '.docx':
        return extract_text_from_docx(file_path)
    if ext == '.txt':
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read().strip()
    raise ValueError(f"Unsupported document type: {ext or 'unknown'}")


def extract_text_from_pdf(file_path: str) -> str:
    if pdfplumber is None:
        raise ValueError("PDF support is not installed on the AI service (pdfplumber missing).")
    try:
        pages = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    pages.append(page_text)
    except Exception as exc:
        raise ValueError("The PDF could not be read. It may be corrupted or password-protected.") from exc
    return "\n".join(pages).strip()


def extract_text_from_docx(file_path: str) -> str:
    if docx is None:
        raise ValueError("DOCX support is not installed on the AI service (python-docx missing).")
    try:
        document = docx.Document(file_path)
    except Exception as exc:
        raise ValueError("The DOCX file could not be read. It may be corrupted.") from exc

    parts = [para.text for para in document.paragraphs]
    # Many resumes keep skills and contact details inside tables
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()
